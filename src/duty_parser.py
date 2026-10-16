#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""值班表解析器：从排课/值班 Excel 中提取指定日期的值班人员名单。

用法:
    python duty_parser.py --file 值班表.xlsx --date 2026-10-16
    python duty_parser.py --file 值班表.xlsx --today   # 自动用今天

输入 Excel 约定（长表格式，一行一条记录，表头自动识别）:
    日期 | 姓名 | (班级/科目/时段 等任意附加列)

日期列兼容: datetime 单元格、2026-10-16、2026/10/16、2026.10.16、10月16日。
若表中没有"日期"列但有"星期"列（周一~周日），则按目标日期的星期匹配。
"""
import argparse
import datetime as dt
import json
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    print(json.dumps({"ok": False, "error": "缺少 openpyxl，请先 pip install openpyxl"},
                     ensure_ascii=False))
    sys.exit(1)

WEEK_MAP = {"周一": 0, "周二": 1, "周三": 2, "周四": 3, "周五": 4, "周六": 5, "周日": 6,
            "星期一": 0, "星期二": 1, "星期三": 2, "星期四": 3, "星期五": 4, "星期六": 5, "星期日": 6}

NAME_KEYS = ("姓名", "值班人", "老师", "教师", "名字")
DATE_KEYS = ("日期", "值班日期", "时间")
WEEK_KEYS = ("星期", "周几", "礼拜")


def _norm_date(v):
    """把各种日期写法归一化为 (yyyy, mm, dd)。"""
    if isinstance(v, dt.datetime):
        return (v.year, v.month, v.day)
    if isinstance(v, dt.date):
        return (v.year, v.month, v.day)
    s = str(v).strip().replace("年", "-").replace("月", "-").replace("日", "")
    s = s.replace("/", "-").replace(".", "-").replace("／", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return tuple(int(x) for x in m.groups())
    m = re.match(r"^(\d{1,2})-(\d{1,2})$", s)  # 无年份：10-16
    if m:
        today = dt.date.today()
        return (today.year, int(m.group(1)), int(m.group(2)))
    return None


def _find_header(rows):
    """在前 10 行里找表头行，返回 (行号, {字段: 列号})。"""
    for i, row in enumerate(rows[:10]):
        cells = [str(c).strip() if c is not None else "" for c in row]
        colmap = {}
        for j, c in enumerate(cells):
            if any(k in c for k in NAME_KEYS):
                colmap["name"] = j
            if any(k in c for k in DATE_KEYS):
                colmap["date"] = j
            if any(k in c for k in WEEK_KEYS):
                colmap["week"] = j
        if "name" in colmap and ("date" in colmap or "week" in colmap):
            return i, colmap
    return None, None


def parse_duty(path, target=None, sheet=None):
    target = target or dt.date.today()
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx, colmap = _find_header(rows)
    if header_idx is None:
        return {"ok": False,
                "error": "未识别到表头：需要包含「姓名」列和「日期/星期」列"}

    people, seen = [], set()
    for row in rows[header_idx + 1:]:
        if row is None or all(c is None for c in row):
            continue
        name = row[colmap["name"]]
        name = str(name).strip() if name is not None else ""
        if not name or "姓名" in name:
            continue
        matched = False
        if "date" in colmap:
            d = _norm_date(row[colmap["date"]])
            matched = d == (target.year, target.month, target.day)
        elif "week" in colmap:
            w = str(row[colmap["week"]] or "").strip()
            matched = WEEK_MAP.get(w, -1) == target.weekday()
        if not matched:
            continue
        extras = {str(rows[header_idx][j]).strip(): ("" if row[j] is None else str(row[j]).strip())
                  for j in range(len(row))
                  if j not in (colmap.get("date"), colmap.get("week"), colmap["name"])
                  and j < len(rows[header_idx]) and rows[header_idx][j]}
        key = (name, tuple(sorted(extras.items())))
        if key not in seen:
            seen.add(key)
            people.append({"name": name, **extras})

    return {"ok": True, "date": target.isoformat(), "weekday": "周" + "一二三四五六日"[target.weekday()],
            "count": len(people), "people": people}


def main():
    ap = argparse.ArgumentParser(description="值班表解析器")
    ap.add_argument("--file", required=True, help="值班/排课表 xlsx 路径")
    ap.add_argument("--date", help="目标日期 YYYY-MM-DD，缺省为今天")
    ap.add_argument("--today", action="store_true", help="使用今天日期")
    ap.add_argument("--sheet", help="工作表名，缺省取第一个")
    args = ap.parse_args()

    if args.date:
        y, m, d = (int(x) for x in args.date.split("-"))
        target = dt.date(y, m, d)
    elif args.today:
        target = dt.date.today()
    else:
        target = dt.date.today()  # 默认今天，值班餐管家的典型场景

    result = parse_duty(args.file, target=target, sheet=args.sheet)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    sys.exit(0 if result.get("ok") else 1)


if __name__ == "__main__":
    main()
